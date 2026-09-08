"""T868 — فلاتر جدول مركبات المزاد، على صفوفٍ حقيقيةٍ في القاعدة.

المعيار I6 حرفياً: لا يُكتفى بحالة 200 — كلُّ اختبارٍ يقرأ الصفوفَ الراجعةَ
والعدّاداتِ المرسومة. وأخطرُ ما يُحرَس هنا فلترةُ الصفحة المعروضة وحدها بدل
المزاد كلّه: الترقيمُ على الخادم (25 صفّاً)، ففلترٌ بعد الترقيم يقول «٣ بلا
صور» والحقيقةُ ٤٠ — ورقمٌ كهذا أسوأ من غيابه.
"""

from __future__ import annotations

from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone

from apps.auctions import engine
from apps.auctions.models import Auction, Vehicle, VehicleImage
from apps.auctions.states import AuctionState, VehicleState
from apps.console import vehicle_filters
from apps.console.auctions import PAGE_SIZE
from apps.core.permissions import Role

pytestmark = pytest.mark.django_db


def make_staff(role: str, phone: str):
    from apps.accounts.models import User

    user = User.objects.create_user(phone=phone, full_name="موظف", password="x")
    user.is_staff = True
    user.console_role = role
    user.save(update_fields=["is_staff", "console_role"])
    return user


@pytest.fixture
def operator(client):
    user = make_staff(Role.OPERATIONS, "966500000101")
    client.force_login(user)
    return user


@pytest.fixture
def auction(db) -> Auction:
    now = timezone.now()
    return Auction.objects.create(
        number=701,
        title="مزاد الفلاتر",
        starts_at=now - timezone.timedelta(hours=1),
        ends_at=now + timezone.timedelta(hours=1),
        state=AuctionState.LIVE,
        deposit_required=Decimal("10000.00"),
    )


def a_car(auction: Auction, lot: int, **extra) -> Vehicle:
    fields = {
        "auction": auction,
        "lot_number": lot,
        "make": "تويوتا",
        "model": "كامري",
        "year": 2020,
        "state": VehicleState.LISTED,
        "reserve_price": Decimal("55000.00"),
    }
    fields.update(extra)
    return Vehicle.objects.create(**fields)


def add_image(vehicle: Vehicle, name: str = "cars/test.jpg") -> VehicleImage:
    # اسمُ ملفٍّ لا ملفّ: العدُّ استعلامٌ لا قراءةَ بايتات.
    return VehicleImage.objects.create(vehicle=vehicle, image=name)


@pytest.fixture
def thirty(auction) -> Auction:
    """ثلاثون مركبة، سبعٌ منها بلا صور — فوق حجم الصفحة (25) عمداً.

    والسبعُ في ذيل الترتيب (لوت 24-30) لا في أوّله: خمسٌ منها خارجَ الصفحة
    الأولى، ففلترٌ يقرأ الصفحةَ وحدها يعيد ٢ لا ٧ — وهذا بالضبط ما يُحرَس.
    """
    for lot in range(1, 31):
        car = a_car(auction, lot)
        if lot <= 23:
            add_image(car, f"cars/test{lot}.jpg")
    return auction


# ---------------------------------------------------------------------------
# ١ — فلترُ «بلا صور» يقرأ المزادَ كلَّه لا الصفحة
# ---------------------------------------------------------------------------


def test_photoless_filter_reads_the_whole_auction(thirty):
    rows = engine.vehicle_rows(thirty)

    assert rows.count() == 30 > PAGE_SIZE

    photoless = vehicle_filters.apply(rows, {"photos": "without"})
    assert photoless.count() == 7
    assert sorted(photoless.values_list("lot_number", flat=True)) == [24, 25, 26, 27, 28, 29, 30]

    tallies = vehicle_filters.counts(rows)
    assert tallies["photos"] == {"with": 23, "without": 7}


def test_photoless_page_shows_seven_not_twenty_five(client, operator, thirty):
    url = reverse("console:auction-detail", args=[thirty.pk])
    response = client.get(url, {"photos": "without"})

    assert response.status_code == 200
    page = response.context["page"]
    assert page.paginator.count == 7
    assert "7 مركبة" in response.content.decode()


def test_unfiltered_first_page_still_paginates(client, operator, thirty):
    url = reverse("console:auction-detail", args=[thirty.pk])
    response = client.get(url)

    assert response.status_code == 200
    page = response.context["page"]
    assert page.paginator.count == 30
    assert len(page.object_list) == PAGE_SIZE
    assert "صفحة 1 من 2" in response.content.decode()


# ---------------------------------------------------------------------------
# ٢ — التطبيع: أرقامٌ عربية-هندية، و«الماركه» تجد «الماركة»
# ---------------------------------------------------------------------------


def test_arabic_indic_digits_find_the_ascii_plate(auction):
    a_car(auction, 1, plate_number="ABC 1234")
    a_car(auction, 2, plate_number="XYZ 5678")

    found = vehicle_filters.apply(engine.vehicle_rows(auction), {"q_plate": "١٢٣٤"})

    assert [car.lot_number for car in found] == [1]


def test_teh_marbuta_matches_teh(auction):
    a_car(auction, 1, make="الماركة")
    a_car(auction, 2, make="هوندا")

    found = vehicle_filters.apply(engine.vehicle_rows(auction), {"q_make": "الماركه"})

    assert [car.lot_number for car in found] == [1]


def test_normalize_folds_digits_and_arabic_letters():
    assert vehicle_filters.normalize("٣١٠") == "310"
    assert vehicle_filters.normalize("۳۱۰") == "310"
    assert vehicle_filters.normalize("الماركه") == vehicle_filters.normalize("الماركة")
    assert vehicle_filters.normalize("أإآ") == "ااا"
    assert vehicle_filters.normalize("علي") == vehicle_filters.normalize("على")
    assert vehicle_filters.normalize("سَيَّارَة") == "سياره"
    assert vehicle_filters.normalize("تويوتا ـ كامري") == "تويوتاكامري"


# ---------------------------------------------------------------------------
# ٣ — فلاترُ معاً: التقاطع لا الاتحاد
# ---------------------------------------------------------------------------


def test_two_filters_intersect(auction):
    a_car(auction, 1, state=VehicleState.DRAFT)  # بلا صور + مسودة — تُراد
    with_image = a_car(auction, 2, state=VehicleState.DRAFT)
    add_image(with_image)  # لها صور + مسودة — لا تُراد
    a_car(auction, 3, state=VehicleState.LISTED)  # بلا صور + معروضة — لا تُراد

    found = vehicle_filters.apply(
        engine.vehicle_rows(auction),
        {"photos": "without", "state": VehicleState.DRAFT},
    )

    assert [car.lot_number for car in found] == [1]


def test_marketing_tabs_read_is_marketing(auction):
    a_car(auction, 1, is_marketing=True)
    a_car(auction, 2, is_marketing=False)

    assert [c.lot_number for c in vehicle_filters.apply(
        engine.vehicle_rows(auction), {"marketing": "1"})] == [1]
    assert [c.lot_number for c in vehicle_filters.apply(
        engine.vehicle_rows(auction), {"marketing": "0"})] == [2]


# ---------------------------------------------------------------------------
# ٤ — التبويبةُ تحفظ البحثَ المكتوب، والعدّاداتُ قبل الترشيح
# ---------------------------------------------------------------------------


def test_a_tab_link_keeps_the_typed_search(auction):
    a_car(auction, 1)

    tabs = vehicle_filters.state(
        {"q_plate": "123", "photos": "without"}, auction
    )

    for tab in tabs["photo_tabs"] + tabs["state_tabs"] + tabs["mkt_tabs"]:
        assert "q_plate=123" in tab["url"], tab["url"]
    without = next(t for t in tabs["photo_tabs"] if t["value"] == "without")
    assert without["active"]


def test_counts_are_computed_before_filtering(auction):
    a_car(auction, 1, state=VehicleState.DRAFT)
    imaged = a_car(auction, 2, state=VehicleState.LISTED)
    add_image(imaged)

    tabs = vehicle_filters.state({"state": VehicleState.DRAFT}, auction)

    photo_counts = {t["value"]: t["count"] for t in tabs["photo_tabs"]}
    assert photo_counts == {"": 2, "without": 1, "with": 1}
    mkt_counts = {t["value"]: t["count"] for t in tabs["mkt_tabs"]}
    assert mkt_counts[""] == 2
    draft = next(t for t in tabs["state_tabs"] if t["value"] == VehicleState.DRAFT)
    assert draft["active"] and draft["count"] == 1
    listed = next(t for t in tabs["state_tabs"] if t["value"] == VehicleState.LISTED)
    assert listed["count"] == 1


def test_zero_counts_still_render(client, operator, auction):
    """«٠ بلا صور» جوابٌ يُقرأ — العدّادُ يُعرَض ولو كان صفراً."""
    imaged = a_car(auction, 1)
    add_image(imaged)

    body = client.get(reverse("console:auction-detail", args=[auction.pk])).content.decode()

    assert "بلا صور" in body
    assert "امسح الكل" not in body


def test_clear_all_appears_only_with_an_active_filter(client, operator, thirty):
    url = reverse("console:auction-detail", args=[thirty.pk])

    assert "امسح الكل" not in client.get(url).content.decode()
    assert "امسح الكل" in client.get(url, {"photos": "without"}).content.decode()


def test_state_tabs_list_only_states_present_in_this_auction(auction):
    a_car(auction, 1, state=VehicleState.DRAFT)
    a_car(auction, 2, state=VehicleState.LISTED)

    tabs = vehicle_filters.state({}, auction)

    assert [t["value"] for t in tabs["state_tabs"]] == [
        "",
        VehicleState.DRAFT,
        VehicleState.LISTED,
    ]
    assert tabs["clear_url"].endswith(
        reverse("console:auction-detail", args=[auction.pk])
    )


def test_an_unknown_state_value_is_ignored(auction):
    a_car(auction, 1)

    found = vehicle_filters.apply(engine.vehicle_rows(auction), {"state": "bogus"})

    assert found.count() == 1


# ---------------------------------------------------------------------------
# ما يغادر الصفحة يحمل الفلتر معه — الثغرتان اللتان تركهما حدُّ الأسطر
# ---------------------------------------------------------------------------


def test_the_next_page_link_keeps_the_filter(client, operator, thirty):
    """«التالي» يبقيك داخل نتيجتك.

    ولولا ذلك لعاد الضاغطُ إلى الجدول كلِّه **وهو يظنّ نفسه داخل المُرشَّح** —
    والصفحةُ الثانية من ثلاثين تبدو كصفحةٍ ثانيةٍ من سبعة لمن لا يعدّ.
    """
    url = reverse("console:auction-detail", args=[thirty.pk])
    # فلترٌ يُبقي أكثر من صفحة: ٢٣ صفّاً «لها صور» وحجمُ الصفحة ٢٥ — فنرشّح
    # بالحالة بدلها لنضمن صفحتين.
    response = client.get(url, {"state": VehicleState.LISTED})
    body = response.content.decode()

    assert "page=2" in body, "لا رابطَ لصفحةٍ ثانية — التجهيزة لم تعد تُرقّم"
    for link in body.split('href="')[1:]:
        href = link.split('"')[0]
        if "page=2" in href:
            assert f"state={VehicleState.LISTED}" in href, (
                f"رابطُ الصفحة الثانية أسقط الفلتر: {href}"
            )
            break


def test_the_export_button_carries_the_filter(client, operator, thirty):
    """الملفُّ يحمل ما تراه الشاشة.

    وتصديرٌ يتجاهل الفلتر يعطي ثلاثين صفّاً والشاشةُ تقول سبعة، ولا شيء في
    الملفّ يقول أيَّهما الصحيح — وهو العطلُ نفسه الذي أُصلح في `0ad74b7` على
    مستوى المزاد، عائداً على مستوى الفلتر.
    """
    url = reverse("console:auction-detail", args=[thirty.pk])
    body = client.get(url, {"photos": "without"}).content.decode()

    exports = [
        link.split('"')[0]
        for link in body.split('href="')[1:]
        if "export" in link.split('"')[0]
    ]
    assert exports, "زرُّ التصدير اختفى"
    assert any("photos=without" in href for href in exports), (
        f"زرُّ التصدير أسقط الفلتر: {exports}"
    )


def test_the_exported_workbook_holds_only_the_filtered_rows(client, operator, thirty):
    """والقياسُ على الملفّ نفسه لا على الرابط: سبعةُ صفوفٍ لا ثلاثون."""
    from io import BytesIO

    from openpyxl import load_workbook

    url = reverse("console:auction-detail", args=[thirty.pk])
    response = client.get(url, {"photos": "without", "export": "xlsx"})

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    # صفٌّ للرأس، ثم سبعةٌ بلا صور.
    assert sheet.max_row == 8, f"الملفّ فيه {sheet.max_row - 1} صفّاً لا ٧"


def test_keep_drops_the_page_but_holds_the_rest():
    """`keep` تُسقط `page` عمداً: الرابطُ ينشئ صفحتَه بنفسه."""
    kept = vehicle_filters.keep({"photos": "without", "page": "3", "q_plate": "أ ب ج"})

    assert "photos=without" in kept
    assert "q_plate" in kept
    assert "page" not in kept
